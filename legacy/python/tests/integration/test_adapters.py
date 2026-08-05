from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook

from conciliacao.adapters import read_bank, read_ledger


def test_ledger_csv_imports_valid_rows_and_reports_invalid_rows(tmp_path):
    path = tmp_path / "RAZAO_SPREAD_MATRIZ_202607.csv"
    path.write_text(
        "Empresa;Filial;Conta contábil;Data;Documento;Histórico;Débito;Crédito;Identificador\n"
        "SPREAD;MATRIZ;1.1.1.01.001;28/07/2026;001;Pagamento;100,00;0;L1\n"
        "SPREAD;MATRIZ;1.1.1.01.001;29/07/2026;002;Recebimento;0;250,50;L2\n"
        "SPREAD;MATRIZ;1.1.1.01.001;31/02/2026;003;Inválido;0;10;L3\n",
        encoding="utf-8",
    )

    result = read_ledger(path, "202607")

    assert [record.value for record in result.records] == [
        Decimal("-100.00"),
        Decimal("250.50"),
    ]
    assert result.records[0].source_line == 2
    assert result.records[0].document == "001"
    assert len(result.errors) == 1
    assert result.errors[0].line == 4
    assert result.total_debit == Decimal("100.00")
    assert result.total_credit == Decimal("250.50")


def test_bank_xlsx_uses_filename_identity_and_keeps_source_line(tmp_path):
    path = tmp_path / "EXTRATO_ITAU_1234_567890_202607.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Data", "Documento", "Histórico", "Valor", "Identificador bancário"])
    sheet.append(["28/07/2026", "001", "Pagamento", "-100,00", "FIT-1"])
    workbook.save(path)

    result = read_bank(path, "202607")

    assert len(result.errors) == 0
    assert result.records[0].bank == "ITAU"
    assert result.records[0].agency == "1234"
    assert result.records[0].bank_account == "567890"
    assert result.records[0].source_id == "FIT-1"
    assert result.records[0].value == Decimal("-100.00")
    assert result.records[0].source_line == 2


def test_protheus_ledger_xlsx_finds_transaction_sheet_and_uses_bank_account_sign(tmp_path):
    path = tmp_path / "RAZAO_SPREAD_01_202602.xlsx"
    workbook = Workbook()
    parameters = workbook.active
    parameters.title = "1-Parametros"
    parameters.append(["Parametros"])
    parameters.append(["Pergunta", "Resposta"])
    transactions = workbook.create_sheet("3-Lançamentos Contábeis")
    transactions.append(["Lançamentos Contábeis"])
    transactions.append(
        [
            "CONTA",
            "DATA",
            "LOTE/SUB/DOC/LINHA",
            "HISTORICO",
            "C/PARTIDA",
            "FILIAL DE ORIGEM",
            "C CUSTO",
            "ITEM CONTA",
            "COD CL VAL",
            "DEBITO",
            "CREDITO",
            "SALDO ATUAL",
        ]
    )
    transactions.append(
        [
            "11102027",
            datetime(2026, 2, 2),
            "008850001000001006",
            "TARIFA BANCARIA",
            "4.2.3.0.1005",
            "01",
            "15",
            "15",
            "15",
            0,
            57.30,
            "2.422.206,07 D",
        ]
    )
    workbook.save(path)

    result = read_ledger(path, "202602")

    assert result.errors == []
    assert len(result.records) == 1
    record = result.records[0]
    assert record.company == "SPREAD"
    assert record.branch == "01"
    assert record.ledger_account == "11102027"
    assert record.value == Decimal("-57.30")
    assert record.source_id == "008850001000001006"
    assert record.source_line == 3


def test_protheus_report_xlsx_infers_account_from_metadata_and_skips_repeated_headers(tmp_path):
    path = tmp_path / "RAZAO_SPREAD_01_202512.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "10-01 - Emissao do Razao Conta"
    headers = [
        "DATA",
        "LOTE/SUB/DOC/LINHA",
        "HISTORICO",
        "C/PARTIDA",
        "FILIAL DE ORIGEM",
        "C CUSTO",
        "ITEM CONTA",
        "COD CL VAL",
        "DEBITO",
        "CREDITO",
        "SALDO ATUAL",
    ]
    sheet.append(headers)
    sheet.append(["CONTA", "DESCRICAO"])
    sheet.append(["1.1.1.0.2", "BANCOS CONTA MOVIMENTOS"])
    sheet.append(["CONTA - 1.1.1.0.2027 - BANCO ITAU", "SALDO ANTERIOR: 2.377.599,34 D"])
    sheet.append(
        [
            datetime(2025, 12, 1),
            "008850001000001005",
            "BAIXA TITULO",
            "1.1.2.0.1001",
            "01",
            "12",
            "12",
            "12",
            23087.50,
            0,
            "2.400.686,84 D",
        ]
    )
    sheet.append(headers)
    sheet.append(
        [
            datetime(2025, 12, 2),
            "008850001000001006",
            "PAGAMENTO",
            "2.1.1.0.1001",
            "01",
            "12",
            "12",
            "12",
            0,
            57.30,
            "2.400.629,54 D",
        ]
    )
    workbook.save(path)

    result = read_ledger(path, "202512")

    assert result.errors == []
    assert [record.ledger_account for record in result.records] == ["11102027", "11102027"]
    assert [record.value for record in result.records] == [
        Decimal("23087.50"),
        Decimal("-57.30"),
    ]
    assert [record.source_line for record in result.records] == [5, 7]


def test_ofx_adapter_reads_sgml_transactions_without_external_service(tmp_path):
    path = tmp_path / "EXTRATO_ITAU_1234_567890_202607.ofx"
    path.write_text(
        "OFXHEADER:100\nDATA:OFXSGML\nVERSION:102\n<OFX><BANKMSGSRSV1><STMTTRNRS>"
        "<STMTRS><BANKACCTFROM><BANKID>341<BRANCHID>1234<ACCTID>567890</BANKACCTFROM>"
        "<BANKTRANLIST><STMTTRN><TRNTYPE>DEBIT<DTPOSTED>20260728120000[-3:BRT]"
        "<TRNAMT>-100.00<FITID>FIT-1<CHECKNUM>001<MEMO>Pagamento</STMTTRN>"
        "</BANKTRANLIST></STMTRS></STMTTRNRS></BANKMSGSRSV1></OFX>",
        encoding="latin-1",
    )

    result = read_bank(path, "202607")

    assert len(result.errors) == 0
    assert len(result.records) == 1
    assert result.records[0].date.isoformat() == "2026-07-28"
    assert result.records[0].value == Decimal("-100.00")
    assert result.records[0].document == "001"
    assert result.records[0].source_line == "FIT-1"


def test_adapter_rejects_missing_required_columns(tmp_path):
    path = tmp_path / "EXTRATO_ITAU_1234_567890_202607.csv"
    path.write_text("Data;Documento\n28/07/2026;001\n", encoding="utf-8")

    result = read_bank(path, "202607")

    assert result.records == []
    assert result.errors[0].code == "COLUNAS_AUSENTES"
    assert "VALOR" in result.errors[0].description
