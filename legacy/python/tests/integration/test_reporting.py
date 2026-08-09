from datetime import UTC, datetime
from decimal import Decimal

from openpyxl import load_workbook

from conciliacao.models import ProcessedFile, TechnicalError
from conciliacao.reporting import ExcelReportGenerator

from ..unit.test_reconciliation import CONFIG, bank, ledger, reconcile


def test_report_contains_required_sheets_and_traceable_summary(tmp_path):
    results = []
    results += reconcile([ledger("L1")], [bank("B1")])
    results += reconcile(
        [ledger("L2", document="VAL", value="-1500")],
        [bank("B2", document="VAL", value="-1450")],
    )
    results += reconcile([ledger("L3", document="ONLY", value="25")], [])
    results += reconcile([], [bank("B3", document="BANK", value="40")])
    processed = [
        ProcessedFile(
            path=tmp_path / "RAZAO_SPREAD_MATRIZ_202607.xlsx",
            file_type="RAZAO",
            period="202607",
            sha256="a" * 64,
            record_count=3,
            company="SPREAD",
            total_debit=Decimal("3000.00"),
            total_credit=Decimal("25.00"),
        )
    ]
    errors = [
        TechnicalError(
            file="EXTRATO_INVALIDO.csv",
            stage="VALIDACAO_REGISTRO",
            code="DADO_INVALIDO",
            description="Data inválida",
            line=8,
            guidance="Corrija a data.",
            occurred_at=datetime(2026, 7, 31, tzinfo=UTC),
        )
    ]
    output = tmp_path / "CONCILIACAO_SPREAD_202607.xlsx"

    ExcelReportGenerator().generate(
        output,
        results=results,
        accounts=[CONFIG],
        processed_files=processed,
        errors=errors,
        executed_at=datetime(2026, 7, 31, 15, 0, tzinfo=UTC),
        configuration_hash="c" * 64,
    )

    workbook = load_workbook(output, data_only=True)
    assert workbook.sheetnames == [
        "Resumo",
        "Conciliados",
        "Divergencias_Detalhadas",
        "Razao_Sem_Extrato",
        "Extrato_Sem_Razao",
        "Ambiguidades",
        "Arquivos_Processados",
        "Erros_Tecnicos",
    ]
    summary_headers = [cell.value for cell in workbook["Resumo"][1]]
    summary = dict(
        zip(summary_headers, [cell.value for cell in workbook["Resumo"][2]], strict=True)
    )
    assert summary["Quantidade conciliada"] == 1
    assert summary["Quantidade de divergências"] == 1
    assert summary["Quantidade sem extrato"] == 1
    assert summary["Quantidade sem razão"] == 1
    assert summary["Diferença de saldo"] == summary["Saldo do razão"] - summary["Saldo do extrato"]
    assert workbook["Conciliados"].max_row == 2
    assert workbook["Divergencias_Detalhadas"].max_row == 2
    assert workbook["Erros_Tecnicos"].max_row == 2
    assert workbook["Resumo"].freeze_panes == "A2"
    assert workbook["Resumo"].auto_filter.ref


def test_report_neutralizes_formula_like_input(tmp_path):
    dangerous_ledger = ledger(document="FORMULA")
    dangerous_bank = bank(document="FORMULA")
    dangerous_ledger = dangerous_ledger.__class__(
        **{
            field: getattr(dangerous_ledger, field)
            for field in dangerous_ledger.__dataclass_fields__
            if field != "history"
        },
        history='=HYPERLINK("https://invalid")',
    )
    output = tmp_path / "safe.xlsx"

    ExcelReportGenerator().generate(
        output,
        results=reconcile([dangerous_ledger], [dangerous_bank]),
        accounts=[CONFIG],
        processed_files=[],
        errors=[],
        executed_at=datetime(2026, 7, 31, tzinfo=UTC),
        configuration_hash="c" * 64,
    )

    workbook = load_workbook(output, data_only=False)
    headers = [cell.value for cell in workbook["Conciliados"][1]]
    history_column = headers.index("Histórico do razão") + 1
    value = workbook["Conciliados"].cell(2, history_column).value
    assert value.startswith("'=")
    assert workbook["Conciliados"].cell(2, history_column).data_type == "s"
