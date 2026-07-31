"""Cria dados totalmente fictícios para uma demonstração local do fluxo."""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook

from conciliacao.processing import initialize_workspace


def create_demo(base_dir: Path, period: str, *, pdf_statement: bool = False) -> None:
    base_dir = Path(base_dir).resolve()
    config_path = initialize_workspace(base_dir)
    workbook = load_workbook(config_path)
    accounts = workbook["Contas"]
    if accounts.max_row == 1:
        accounts.append(
            [
                "DEMO",
                "MATRIZ",
                "1.1.1.01.001",
                "Banco Demonstração",
                "ITAU" if pdf_statement else "BANCO",
                "1234",
                "567890",
                "CSV,XLSX,OFX,PDF",
                2,
                "0,01",
                "DOC_VALOR_DATA",
                "SIM",
                "Usuário Demo",
            ]
        )
        workbook.save(config_path)
    year, month = period[:4], period[4:]
    ledger = base_dir / "01_Entrada" / "Razao" / f"RAZAO_DEMO_MATRIZ_{period}.csv"
    bank = "ITAU" if pdf_statement else "BANCO"
    suffix = "pdf" if pdf_statement else "csv"
    statement = (
        base_dir / "01_Entrada" / "Extratos" / f"EXTRATO_{bank}_1234_567890_{period}.{suffix}"
    )
    if ledger.exists() or statement.exists():
        raise FileExistsError("A entrada demo já existe; use uma pasta vazia.")
    ledger.write_text(
        "Empresa;Filial;Conta contábil;Data;Documento;Histórico;Débito;Crédito;Identificador\n"
        f"DEMO;MATRIZ;1.1.1.01.001;05/{month}/{year};001;PAGAMENTO FICTICIO;100;0;L1\n"
        f"DEMO;MATRIZ;1.1.1.01.001;06/{month}/{year};002;RECEBIMENTO FICTICIO;0;250;L2\n",
        encoding="utf-8",
    )
    if pdf_statement:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas

        document = canvas.Canvas(str(statement), pagesize=A4)
        lines = [
            "BANCO ITAU - EXTRATO CONTA CORRENTE",
            "TITULAR: EMPRESA DEMONSTRACAO",
            "AGENCIA: 1234 CONTA: 567890",
            f"PERIODO: 01/{month}/{year} A 31/{month}/{year}",
            "DATA DESCRICAO DOCUMENTO VALOR SALDO",
            "SALDO ANTERIOR 1.000,00 C",
            f"05/{month} PAGAMENTO FICTICIO 001 100,00 D 900,00 C",
            f"06/{month} RECEBIMENTO FICTICIO 002 250,00 C 1.150,00 C",
            f"07/{month} TARIFA FICTICIA 003 40,00 D 1.110,00 C",
            "SALDO FINAL 1.110,00 C",
        ]
        y = A4[1] - 55
        document.setFont("Helvetica", 10)
        for line in lines:
            document.drawString(45, y, line)
            y -= 18
        document.save()
    else:
        statement.write_text(
            "Data;Documento;Histórico;Valor;Identificador bancário\n"
            f"05/{month}/{year};001;PAGAMENTO FICTICIO;-100;B1\n"
            f"06/{month}/{year};002;RECEBIMENTO FICTICIO;250;B2\n"
            f"07/{month}/{year};003;TARIFA FICTICIA;-40;B3\n",
            encoding="utf-8",
        )


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-dir", type=Path, required=True)
    parser.add_argument("--periodo", default="202607")
    parser.add_argument("--pdf", action="store_true", help="gerar o extrato demo em PDF")
    args = parser.parse_args()
    create_demo(args.base_dir, args.periodo, pdf_statement=args.pdf)
    print(f"Workspace demo criado em {args.base_dir.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
