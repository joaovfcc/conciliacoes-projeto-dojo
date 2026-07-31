"""Adaptadores reais para razão CSV/XLSX e extrato CSV/XLSX/OFX/PDF."""

from __future__ import annotations

import csv
import re
from datetime import UTC, datetime
from decimal import Decimal
from hashlib import sha1
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from conciliacao.discovery import parse_input_filename
from conciliacao.errors import DataValidationError, FileValidationError
from conciliacao.models import BankRecord, LedgerRecord, ReadResult, TechnicalError
from conciliacao.normalization import (
    normalize_code,
    normalize_document,
    normalize_header,
    parse_date,
    parse_money,
)

ALIASES = {
    "EMPRESA": {"EMPRESA", "COD_EMPRESA"},
    "FILIAL": {"FILIAL", "COD_FILIAL"},
    "CONTA_CONTABIL": {"CONTA_CONTABIL", "CONTA", "CT2_DEBITO", "CT2_CREDIT"},
    "DATA": {"DATA", "DT_LANCAMENTO", "DATA_LANCAMENTO", "DTPOSTED"},
    "DOCUMENTO": {"DOCUMENTO", "DOC", "NUM_DOCUMENTO", "CHECKNUM", "REFNUM"},
    "HISTORICO": {"HISTORICO", "DESCRICAO", "MEMO"},
    "DEBITO": {"DEBITO", "VALOR_DEBITO"},
    "CREDITO": {"CREDITO", "VALOR_CREDITO"},
    "VALOR": {"VALOR", "VALOR_LIQUIDO", "TRNAMT"},
    "IDENTIFICADOR": {"IDENTIFICADOR", "ID", "ID_PROTHEUS"},
    "IDENTIFICADOR_BANCARIO": {"IDENTIFICADOR_BANCARIO", "FITID", "ID_BANCARIO"},
}


def read_ledger(path: Path, period: str) -> ReadResult:
    path = Path(path)
    parse_input_filename(path)
    result = ReadResult()
    try:
        rows = _tabular_rows(path)
        headers = set(rows[0][1]) if rows else set()
        missing = {"EMPRESA", "FILIAL", "CONTA_CONTABIL", "DATA"} - headers
        if "VALOR" not in headers and not {"DEBITO", "CREDITO"}.issubset(headers):
            missing.add("VALOR ou DEBITO+CREDITO")
        if missing:
            return _missing_columns(path, missing)
        for line, row in rows:
            try:
                record_date = parse_date(row.get("DATA"))
                _validate_period(record_date, period)
                company = _required_code(row, "EMPRESA")
                branch = _required_code(row, "FILIAL")
                account = _required_code(row, "CONTA_CONTABIL")
                if "VALOR" in row and row.get("VALOR") not in (None, ""):
                    value = parse_money(row["VALOR"])
                    debit = abs(value) if value < 0 else Decimal("0.00")
                    credit = value if value > 0 else Decimal("0.00")
                else:
                    debit = _optional_money(row.get("DEBITO"))
                    credit = _optional_money(row.get("CREDITO"))
                    value = credit - debit
                source_id = _source_identifier(row.get("IDENTIFICADOR")) or _generated_id(
                    "RAZ", path.name, line
                )
                result.records.append(
                    LedgerRecord(
                        record_id=_generated_id("LR", path.name, line),
                        company=company,
                        branch=branch,
                        ledger_account=account,
                        date=record_date,
                        document=normalize_document(row.get("DOCUMENTO")),
                        history=str(row.get("HISTORICO") or "").strip(),
                        value=value,
                        nature=_nature(value),
                        source_id=source_id,
                        source_file=path.name,
                        source_line=line,
                    )
                )
                result.total_debit += debit
                result.total_credit += credit
            except DataValidationError as exc:
                result.errors.append(_row_error(path, line, exc))
    except (OSError, ValueError, FileValidationError) as exc:
        result.errors.append(_file_error(path, "LEITURA_RAZAO", exc))
    return result


def read_bank(path: Path, period: str) -> ReadResult:
    path = Path(path)
    descriptor = parse_input_filename(path)
    if descriptor.file_type != "EXTRATO":
        return ReadResult(
            errors=[_file_error(path, "LEITURA_EXTRATO", "Tipo de arquivo incorreto.")]
        )
    if descriptor.extension == "OFX":
        return _read_ofx(path, period, descriptor.bank, descriptor.agency, descriptor.account)
    if descriptor.extension == "PDF":
        from conciliacao.pdf_statements.service import read_pdf_statement

        return read_pdf_statement(path, period)

    result = ReadResult()
    try:
        rows = _tabular_rows(path)
        headers = set(rows[0][1]) if rows else set()
        if missing := {"DATA", "VALOR"} - headers:
            return _missing_columns(path, missing)
        for line, row in rows:
            try:
                record_date = parse_date(row.get("DATA"))
                _validate_period(record_date, period)
                value = parse_money(row.get("VALOR"))
                source_id = _source_identifier(row.get("IDENTIFICADOR_BANCARIO")) or _generated_id(
                    "BAN", path.name, line
                )
                result.records.append(
                    BankRecord(
                        record_id=_generated_id("BR", path.name, line),
                        bank=descriptor.bank,
                        agency=descriptor.agency,
                        bank_account=descriptor.account,
                        date=record_date,
                        document=normalize_document(row.get("DOCUMENTO")),
                        history=str(row.get("HISTORICO") or "").strip(),
                        value=value,
                        nature=_nature(value),
                        source_id=source_id,
                        source_file=path.name,
                        source_line=line,
                    )
                )
                result.total_debit += abs(value) if value < 0 else Decimal("0.00")
                result.total_credit += value if value > 0 else Decimal("0.00")
            except DataValidationError as exc:
                result.errors.append(_row_error(path, line, exc))
    except (OSError, ValueError, FileValidationError) as exc:
        result.errors.append(_file_error(path, "LEITURA_EXTRATO", exc))
    return result


def _tabular_rows(path: Path) -> list[tuple[int, dict[str, Any]]]:
    if path.suffix.lower() == ".csv":
        raw_rows = _csv_rows(path)
    elif path.suffix.lower() == ".xlsx":
        raw_rows = _xlsx_rows(path)
    else:
        raise FileValidationError(f'A extensão "{path.suffix}" não é tabular ou autorizada.')
    if not raw_rows:
        raise FileValidationError(f'O arquivo "{path.name}" não contém registros.')
    headers = [_canonical_header(value) for value in raw_rows[0]]
    if not any(headers):
        raise FileValidationError(f'O arquivo "{path.name}" não possui cabeçalho válido.')
    return [
        (line, dict(zip(headers, values, strict=False)))
        for line, values in enumerate(raw_rows[1:], start=2)
        if any(value not in (None, "") for value in values)
    ]


def _csv_rows(path: Path) -> list[list[str]]:
    content = None
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            content = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    if content is None:
        raise FileValidationError(f'Não foi possível identificar a codificação de "{path.name}".')
    try:
        dialect = csv.Sniffer().sniff(content[:4096], delimiters=";,\t|")
        return list(csv.reader(content.splitlines(), dialect))
    except csv.Error as exc:
        raise FileValidationError(f'O CSV "{path.name}" possui estrutura inválida.') from exc


def _xlsx_rows(path: Path) -> list[list[Any]]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
        return rows
    except Exception as exc:
        raise FileValidationError(f'O XLSX "{path.name}" está corrompido ou bloqueado.') from exc


def _read_ofx(path: Path, period: str, bank: str, agency: str, account: str) -> ReadResult:
    result = ReadResult()
    try:
        content = path.read_text(encoding="latin-1", errors="replace")
        blocks = re.findall(
            r"<STMTTRN>(.*?)(?:</STMTTRN>|(?=<STMTTRN>|</BANKTRANLIST>))", content, re.I | re.S
        )
        if not blocks:
            raise FileValidationError(f'O OFX "{path.name}" não possui movimentações STMTTRN.')
        for position, block in enumerate(blocks, start=1):
            try:
                date_text = _ofx_field(block, "DTPOSTED", required=True)[:8]
                record_date = parse_date(date_text)
                _validate_period(record_date, period)
                value = parse_money(_ofx_field(block, "TRNAMT", required=True))
                fitid = _source_identifier(_ofx_field(block, "FITID"))
                line: int | str = fitid or position
                result.records.append(
                    BankRecord(
                        record_id=_generated_id("BR", path.name, line),
                        bank=bank,
                        agency=agency,
                        bank_account=account,
                        date=record_date,
                        document=normalize_document(
                            _ofx_field(block, "CHECKNUM") or _ofx_field(block, "REFNUM")
                        ),
                        history=_ofx_field(block, "MEMO") or _ofx_field(block, "NAME"),
                        value=value,
                        nature=_nature(value),
                        source_id=fitid or _generated_id("OFX", path.name, position),
                        source_file=path.name,
                        source_line=line,
                    )
                )
                result.total_debit += abs(value) if value < 0 else Decimal("0.00")
                result.total_credit += value if value > 0 else Decimal("0.00")
            except DataValidationError as exc:
                result.errors.append(_row_error(path, position, exc))
    except (OSError, FileValidationError) as exc:
        result.errors.append(_file_error(path, "LEITURA_OFX", exc))
    return result


def _ofx_field(block: str, tag: str, *, required: bool = False) -> str:
    match = re.search(rf"<{tag}>([^<\r\n]+)", block, re.I)
    value = match.group(1).strip() if match else ""
    if required and not value:
        raise DataValidationError(f"O campo OFX {tag} está ausente.")
    return value


def _canonical_header(value: Any) -> str:
    normalized = normalize_header(value)
    for canonical, aliases in ALIASES.items():
        if normalized in aliases:
            return canonical
    return normalized


def _validate_period(value, period: str) -> None:
    if value.strftime("%Y%m") != period:
        raise DataValidationError(
            f'A data "{value:%d/%m/%Y}" não pertence ao período solicitado {period}.'
        )


def _required_code(row: dict[str, Any], field: str) -> str:
    value = normalize_code(row.get(field))
    if not value:
        raise DataValidationError(f"O campo obrigatório {field} está ausente.")
    return value


def _optional_money(value: Any) -> Decimal:
    return Decimal("0.00") if value in (None, "") else parse_money(value)


def _nature(value: Decimal) -> str:
    if value < 0:
        return "DEBITO"
    if value > 0:
        return "CREDITO"
    return "NEUTRO"


def _generated_id(prefix: str, filename: str, line: int | str) -> str:
    digest = sha1(f"{filename}|{line}".encode(), usedforsecurity=False).hexdigest()[:12].upper()
    return f"{prefix}-{digest}"


def _source_identifier(value: Any) -> str:
    return str(value or "").strip()


def _missing_columns(path: Path, missing: set[str]) -> ReadResult:
    names = ", ".join(sorted(missing))
    return ReadResult(
        errors=[
            TechnicalError(
                file=path.name,
                stage="VALIDACAO_ESTRUTURA",
                code="COLUNAS_AUSENTES",
                description=f"Colunas obrigatórias ausentes: {names}.",
                line=1,
                guidance="Use o layout documentado e mantenha os cabeçalhos obrigatórios.",
                occurred_at=datetime.now(UTC),
            )
        ]
    )


def _row_error(path: Path, line: int | str, error: Exception) -> TechnicalError:
    return TechnicalError(
        file=path.name,
        stage="VALIDACAO_REGISTRO",
        code=getattr(error, "code", "DADO_INVALIDO"),
        description=str(error),
        line=line,
        guidance="Corrija o registro indicado e processe novamente.",
        occurred_at=datetime.now(UTC),
    )


def _file_error(path: Path, stage: str, error: Exception | str) -> TechnicalError:
    return TechnicalError(
        file=path.name,
        stage=stage,
        code=getattr(error, "code", "ERRO_DE_LEITURA"),
        description=str(error),
        line=None,
        guidance="Verifique formato, conteúdo e permissão do arquivo.",
        occurred_at=datetime.now(UTC),
    )
