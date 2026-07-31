"""Planilha operacional de contas, layouts e parâmetros."""

from __future__ import annotations

import hashlib
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from conciliacao.errors import ConfigurationError
from conciliacao.models import AccountConfig, ApplicationConfig
from conciliacao.normalization import normalize_code, normalize_header, parse_money

ACCOUNT_HEADERS = [
    "Empresa",
    "Filial",
    "Conta contábil",
    "Descrição",
    "Banco",
    "Agência",
    "Conta bancária",
    "Formato aceito",
    "Tolerância de data",
    "Tolerância de valor",
    "Regra principal",
    "Ativa",
    "Responsável",
]

LAYOUT_HEADERS = [
    "Código do layout",
    "Tipo",
    "Coluna de data",
    "Coluna de documento",
    "Coluna de histórico",
    "Coluna de débito",
    "Coluna de crédito",
    "Coluna de valor",
    "Separador CSV",
    "Codificação",
]

PARAMETERS = [
    ("PERIODO_PADRAO", ""),
    ("TOLERANCIA_GERAL_VALOR", "0,01"),
    ("TOLERANCIA_GERAL_DATA", "2"),
    ("PERMITIR_REPROCESSAMENTO", "NAO"),
]


def create_configuration_template(path: Path) -> Path:
    """Cria template legível, sem dados reais nem segredos."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    accounts = workbook.active
    accounts.title = "Contas"
    accounts.append(ACCOUNT_HEADERS)
    layouts = workbook.create_sheet("Layouts")
    layouts.append(LAYOUT_HEADERS)
    parameters = workbook.create_sheet("Parametros")
    parameters.append(["Parametro", "Valor"])
    for item in PARAMETERS:
        parameters.append(item)
    instructions = workbook.create_sheet("Instrucoes")
    instructions.append(["CONFIGURAÇÃO DE CONCILIAÇÃO BANCÁRIA"])
    instructions.append(["Preencha uma linha na aba Contas para cada relacionamento ativo."])
    instructions.append(["Feche o Excel antes de executar a aplicação."])
    instructions.append(["Não inclua senhas, tokens ou credenciais neste arquivo."])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for sheet in (accounts, layouts, parameters):
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(
                35, max(14, max(len(str(cell.value or "")) for cell in column) + 2)
            )

    parameter_table = Table(displayName="ParametrosConfig", ref=f"A1:B{parameters.max_row}")
    parameter_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
    )
    parameters.add_table(parameter_table)
    workbook.save(path)
    return path


def load_configuration(path: Path) -> ApplicationConfig:
    path = Path(path)
    if not path.is_file():
        raise ConfigurationError(f'A configuração obrigatória não foi encontrada em "{path}".')
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ConfigurationError(f'Não foi possível abrir a configuração "{path.name}".') from exc
    required_sheets = {"Contas", "Parametros"}
    if missing_sheets := required_sheets.difference(workbook.sheetnames):
        names = ", ".join(sorted(missing_sheets))
        raise ConfigurationError(f"A configuração não contém as abas obrigatórias: {names}.")

    account_rows = _sheet_rows(workbook["Contas"])
    accounts: list[AccountConfig] = []
    for line_number, row in account_rows:
        if not any(value not in (None, "") for value in row.values()):
            continue
        try:
            accounts.append(_account_from_row(row))
        except (ConfigurationError, ValueError) as exc:
            raise ConfigurationError(f"Aba Contas, linha {line_number}: {exc}") from exc
    if not accounts:
        raise ConfigurationError("A aba Contas não possui nenhum relacionamento configurado.")
    ledger_keys = [
        (item.company, item.branch, item.ledger_account) for item in accounts if item.active
    ]
    bank_keys = [item.bank_key for item in accounts if item.active]
    if len(ledger_keys) != len(set(ledger_keys)) or len(bank_keys) != len(set(bank_keys)):
        raise ConfigurationError(
            "Existe relacionamento ativo duplicado de conta contábil ou conta bancária."
        )

    parameters = {
        normalize_header(row.get("PARAMETRO")): str(row.get("VALOR") or "").strip()
        for _, row in _sheet_rows(workbook["Parametros"])
        if row.get("PARAMETRO") not in (None, "")
    }
    layouts = (
        tuple(
            {key: str(value or "").strip() for key, value in row.items()}
            for _, row in _sheet_rows(workbook["Layouts"])
            if any(value not in (None, "") for value in row.values())
        )
        if "Layouts" in workbook.sheetnames
        else ()
    )
    workbook.close()
    return ApplicationConfig(
        accounts=tuple(accounts),
        parameters=parameters,
        layouts=layouts,
        sha256=_sha256(path),
    )


def _sheet_rows(sheet):
    iterator = sheet.iter_rows(values_only=True)
    try:
        headers = [normalize_header(value) for value in next(iterator)]
    except StopIteration:
        return []
    return [
        (line_number, dict(zip(headers, values, strict=False)))
        for line_number, values in enumerate(iterator, start=2)
    ]


def _required(row: dict[str, object], name: str) -> object:
    value = row.get(name)
    if value in (None, ""):
        raise ConfigurationError(f"o campo {name} é obrigatório.")
    return value


def _account_from_row(row: dict[str, object]) -> AccountConfig:
    active_text = normalize_code(_required(row, "ATIVA"))
    if active_text not in {"SIM", "NAO", "NÃO"}:
        raise ConfigurationError('o campo ATIVA deve ser "SIM" ou "NAO".')
    try:
        date_tolerance = int(_required(row, "TOLERANCIA_DE_DATA"))
    except (TypeError, ValueError) as exc:
        raise ConfigurationError("a tolerância de data deve ser um inteiro.") from exc
    if date_tolerance < 0:
        raise ConfigurationError("a tolerância de data não pode ser negativa.")
    return AccountConfig(
        company=normalize_code(_required(row, "EMPRESA")),
        branch=normalize_code(_required(row, "FILIAL")),
        ledger_account=normalize_code(_required(row, "CONTA_CONTABIL")),
        description=str(_required(row, "DESCRICAO")).strip(),
        bank=normalize_code(_required(row, "BANCO")),
        agency=normalize_code(_required(row, "AGENCIA")),
        bank_account=normalize_code(_required(row, "CONTA_BANCARIA")),
        accepted_format=normalize_code(_required(row, "FORMATO_ACEITO")),
        date_tolerance=date_tolerance,
        value_tolerance=parse_money(_required(row, "TOLERANCIA_DE_VALOR")),
        primary_rule=normalize_code(_required(row, "REGRA_PRINCIPAL")),
        active=active_text == "SIM",
        responsible=str(row.get("RESPONSAVEL") or "").strip(),
    )


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
