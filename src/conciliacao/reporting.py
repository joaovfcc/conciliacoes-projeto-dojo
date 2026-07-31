"""Geração do workbook de conciliação e evidências."""

from __future__ import annotations

from collections.abc import Iterable, Sequence
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from conciliacao.models import (
    AccountConfig,
    BankRecord,
    LedgerRecord,
    ProcessedFile,
    ReconciliationResult,
    Status,
    TechnicalError,
)

RECONCILED = {
    Status.CONCILIADO_EXATO,
    Status.CONCILIADO_TOLERANCIA_DATA,
    Status.CONCILIADO_VALOR_DATA,
}
DIVERGENCES = {Status.DIVERGENCIA_VALOR, Status.DIVERGENCIA_DATA}


class ExcelReportGenerator:
    """Cria o output principal com números derivados dos mesmos detalhes."""

    def generate(
        self,
        output_path: Path,
        *,
        results: Sequence[ReconciliationResult],
        accounts: Sequence[AccountConfig],
        processed_files: Sequence[ProcessedFile],
        errors: Sequence[TechnicalError],
        executed_at: datetime,
        configuration_hash: str,
    ) -> Path:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        workbook.remove(workbook.active)
        workbook.properties.title = "Conciliação bancária e evidências"
        workbook.properties.subject = "Resultado auditável de conciliação"
        workbook.properties.creator = "Automação de Conciliação Bancária"
        workbook.properties.keywords = f"config-sha256:{configuration_hash}"

        self._write_sheet(
            workbook,
            "Resumo",
            SUMMARY_HEADERS,
            self._summary_rows(results, accounts, executed_at),
        )
        self._write_sheet(
            workbook,
            "Conciliados",
            MATCH_HEADERS,
            [self._matched_row(item) for item in results if item.status in RECONCILED],
        )
        self._write_sheet(
            workbook,
            "Divergencias_Detalhadas",
            DIVERGENCE_HEADERS,
            [
                self._divergence_row(item, accounts)
                for item in results
                if item.status in DIVERGENCES
            ],
        )
        self._write_sheet(
            workbook,
            "Razao_Sem_Extrato",
            LEDGER_HEADERS,
            [self._ledger_row(item) for item in results if item.status is Status.RAZAO_SEM_EXTRATO],
        )
        self._write_sheet(
            workbook,
            "Extrato_Sem_Razao",
            BANK_HEADERS,
            [self._bank_row(item) for item in results if item.status is Status.EXTRATO_SEM_RAZAO],
        )
        self._write_sheet(
            workbook,
            "Ambiguidades",
            AMBIGUITY_HEADERS,
            self._ambiguity_rows(results),
        )
        self._write_sheet(
            workbook,
            "Arquivos_Processados",
            FILE_HEADERS,
            [self._file_row(item, executed_at) for item in processed_files],
        )
        self._write_sheet(
            workbook,
            "Erros_Tecnicos",
            ERROR_HEADERS,
            [self._error_row(item) for item in errors],
        )
        self._add_analysis_validation(workbook["Divergencias_Detalhadas"])
        self._add_summary_formatting(workbook["Resumo"])
        workbook.save(output_path)
        return output_path

    def _summary_rows(
        self,
        results: Sequence[ReconciliationResult],
        accounts: Sequence[AccountConfig],
        executed_at: datetime,
    ) -> list[list[Any]]:
        rows: list[list[Any]] = []
        for account in accounts:
            account_results = [item for item in results if _belongs_to(item, account)]
            ledgers = _unique_ledgers(account_results)
            banks = _unique_banks(account_results)
            reconciled = [item for item in account_results if item.status in RECONCILED]
            divergences = [item for item in account_results if item.status in DIVERGENCES]
            without_bank = [
                item for item in account_results if item.status is Status.RAZAO_SEM_EXTRATO
            ]
            without_ledger = [
                item for item in account_results if item.status is Status.EXTRATO_SEM_RAZAO
            ]
            ledger_balance = sum((item.value for item in ledgers), Decimal("0.00"))
            bank_balance = sum((item.value for item in banks), Decimal("0.00"))
            exception_count = len(account_results) - len(reconciled)
            rows.append(
                [
                    account.company,
                    account.branch,
                    account.ledger_account,
                    account.bank,
                    account.agency,
                    account.bank_account,
                    ledger_balance,
                    bank_balance,
                    ledger_balance - bank_balance,
                    len(reconciled),
                    sum((abs(item.ledger.value) for item in reconciled), Decimal("0.00")),
                    len(divergences),
                    sum(
                        (abs(item.value_difference or Decimal("0.00")) for item in divergences),
                        Decimal("0.00"),
                    ),
                    len(without_bank),
                    len(without_ledger),
                    "CONCILIADA" if exception_count == 0 else "COM_DIVERGENCIAS",
                    executed_at,
                ]
            )
        return rows

    @staticmethod
    def _matched_row(item: ReconciliationResult) -> list[Any]:
        ledger = item.ledger
        bank = item.bank
        assert ledger is not None and bank is not None
        return [
            item.result_id,
            ledger.ledger_account,
            f"{bank.bank}/{bank.agency}/{bank.bank_account}",
            ledger.date,
            bank.date,
            ledger.document,
            bank.document,
            ledger.history,
            bank.history,
            ledger.value,
            bank.value,
            item.value_difference,
            item.rule,
            ledger.source_file,
            ledger.source_line,
            bank.source_file,
            bank.source_line,
            bank.source_page,
            bank.extraction_method,
            bank.parser_name,
            bank.confidence,
            bank.source_text,
        ]

    @staticmethod
    def _divergence_row(item: ReconciliationResult, accounts: Sequence[AccountConfig]) -> list[Any]:
        ledger = item.ledger
        bank = item.bank
        assert ledger is not None and bank is not None
        account = next((config for config in accounts if _belongs_to(item, config)), None)
        return [
            item.result_id,
            item.status.value,
            ledger.company,
            ledger.branch,
            ledger.ledger_account,
            bank.bank,
            bank.agency,
            bank.bank_account,
            ledger.date,
            ledger.document,
            ledger.history,
            ledger.value,
            ledger.source_file,
            ledger.source_line,
            bank.date,
            bank.document,
            bank.history,
            bank.value,
            bank.source_file,
            bank.source_line,
            bank.source_page,
            bank.extraction_method,
            bank.parser_name,
            bank.confidence,
            bank.source_text,
            item.value_difference,
            item.date_difference,
            item.tolerance_days,
            item.rule,
            item.reason,
            "PENDENTE",
            "",
            "",
            account.responsible if account else "",
            "",
            "",
            "",
        ]

    @staticmethod
    def _ledger_row(item: ReconciliationResult) -> list[Any]:
        ledger = item.ledger
        assert ledger is not None
        return [
            item.result_id,
            item.status.value,
            ledger.company,
            ledger.branch,
            ledger.ledger_account,
            ledger.date,
            ledger.document,
            ledger.history,
            ledger.value,
            ledger.nature,
            ledger.source_id,
            ledger.source_file,
            ledger.source_line,
            item.rule,
            item.reason,
        ]

    @staticmethod
    def _bank_row(item: ReconciliationResult) -> list[Any]:
        bank = item.bank
        assert bank is not None
        return [
            item.result_id,
            item.status.value,
            bank.bank,
            bank.agency,
            bank.bank_account,
            bank.date,
            bank.document,
            bank.history,
            bank.value,
            bank.nature,
            bank.source_id,
            bank.source_file,
            bank.source_line,
            bank.source_page,
            bank.extraction_method,
            bank.parser_name,
            bank.confidence,
            bank.source_text,
            item.rule,
            item.reason,
        ]

    @staticmethod
    def _ambiguity_rows(results: Sequence[ReconciliationResult]) -> list[list[Any]]:
        rows: list[list[Any]] = []
        for item in results:
            if item.status is Status.MULTIPLAS_CORRESPONDENCIAS:
                ledger = item.ledger
                assert ledger is not None
                for candidate in item.candidates:
                    rows.append(
                        [
                            item.result_id,
                            item.status.value,
                            ledger.record_id,
                            ledger.date,
                            ledger.document,
                            ledger.value,
                            ledger.source_file,
                            ledger.source_line,
                            candidate.record_id,
                            candidate.date,
                            candidate.document,
                            candidate.value,
                            candidate.source_file,
                            candidate.source_line,
                            candidate.source_page,
                            candidate.source_text,
                            item.rule,
                            item.reason,
                        ]
                    )
            elif item.status is Status.POSSIVEL_DUPLICIDADE:
                ledger, bank = item.ledger, item.bank
                source = ledger or bank
                assert source is not None
                rows.append(
                    [
                        item.result_id,
                        item.status.value,
                        source.record_id,
                        source.date,
                        source.document,
                        source.value,
                        source.source_file,
                        source.source_line,
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        getattr(source, "source_page", None),
                        getattr(source, "source_text", ""),
                        item.rule,
                        item.reason,
                    ]
                )
        return rows

    @staticmethod
    def _file_row(item: ProcessedFile, executed_at: datetime) -> list[Any]:
        return [
            item.path.name,
            item.file_type,
            item.company,
            item.bank,
            item.agency,
            item.account,
            item.period,
            item.record_count,
            item.total_debit,
            item.total_credit,
            item.page_count,
            item.extraction_method,
            item.valid_record_count,
            item.rejected_record_count,
            item.warning_count,
            item.duration_seconds,
            item.sha256,
            executed_at,
            item.status,
        ]

    @staticmethod
    def _error_row(item: TechnicalError) -> list[Any]:
        return [
            item.file,
            item.stage,
            item.code,
            item.description,
            item.line,
            item.guidance,
            item.occurred_at,
        ]

    def _write_sheet(
        self,
        workbook: Workbook,
        title: str,
        headers: Sequence[str],
        rows: Iterable[Sequence[Any]],
    ) -> None:
        sheet = workbook.create_sheet(title)
        sheet.append(list(headers))
        for row in rows:
            sheet.append([_excel_value(value) for value in row])
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 34
        for column in sheet.columns:
            values = [str(cell.value or "") for cell in column[: min(len(column), 200)]]
            width = min(45, max(12, max(map(len, values), default=10) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, float):
                    cell.number_format = "#,##0.00;[Red]-#,##0.00"
                elif isinstance(cell.value, datetime):
                    cell.number_format = "dd/mm/yyyy hh:mm:ss"
                elif isinstance(cell.value, date):
                    cell.number_format = "dd/mm/yyyy"
        if sheet.max_row >= 2:
            table = Table(
                displayName=f"Tabela{title.replace('_', '')}",
                ref=f"A1:{sheet.cell(1, len(headers)).column_letter}{sheet.max_row}",
            )
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
            )
            sheet.add_table(table)

    @staticmethod
    def _add_analysis_validation(sheet) -> None:
        column = DIVERGENCE_HEADERS.index("Status da análise") + 1
        validation = DataValidation(
            type="list",
            formula1='"PENDENTE,EM_ANALISE,JUSTIFICADO,REGULARIZADO"',
            allow_blank=False,
        )
        sheet.add_data_validation(validation)
        validation.add(
            f"{sheet.cell(1, column).column_letter}2:{sheet.cell(1, column).column_letter}1048576"
        )

    @staticmethod
    def _add_summary_formatting(sheet) -> None:
        status_column = SUMMARY_HEADERS.index("Status da conta") + 1
        letter = sheet.cell(1, status_column).column_letter
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        sheet.conditional_formatting.add(
            f"{letter}2:{letter}{max(sheet.max_row, 2)}",
            CellIsRule(operator="notEqual", formula=['"CONCILIADA"'], fill=red_fill),
        )


def _excel_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _belongs_to(item: ReconciliationResult, account: AccountConfig) -> bool:
    if item.ledger and (
        item.ledger.company,
        item.ledger.branch,
        item.ledger.ledger_account,
    ) != (account.company, account.branch, account.ledger_account):
        return False
    if item.bank and (item.bank.bank, item.bank.agency, item.bank.bank_account) != account.bank_key:
        return False
    return bool(item.ledger or item.bank)


def _unique_ledgers(results: Sequence[ReconciliationResult]) -> list[LedgerRecord]:
    records: dict[str, LedgerRecord] = {}
    for item in results:
        if item.ledger:
            records[item.ledger.record_id] = item.ledger
    return list(records.values())


def _unique_banks(results: Sequence[ReconciliationResult]) -> list[BankRecord]:
    records: dict[str, BankRecord] = {}
    for item in results:
        if item.bank:
            records[item.bank.record_id] = item.bank
        for candidate in item.candidates:
            records[candidate.record_id] = candidate
    return list(records.values())


SUMMARY_HEADERS = [
    "Empresa",
    "Filial",
    "Conta contábil",
    "Banco",
    "Agência",
    "Conta bancária",
    "Saldo do razão",
    "Saldo do extrato",
    "Diferença de saldo",
    "Quantidade conciliada",
    "Valor conciliado",
    "Quantidade de divergências",
    "Valor das divergências",
    "Quantidade sem extrato",
    "Quantidade sem razão",
    "Status da conta",
    "Data e hora da execução",
]

MATCH_HEADERS = [
    "ID da conciliação",
    "Conta contábil",
    "Banco/agência/conta",
    "Data do razão",
    "Data do extrato",
    "Documento do razão",
    "Documento do extrato",
    "Histórico do razão",
    "Histórico do extrato",
    "Valor do razão",
    "Valor do extrato",
    "Diferença",
    "Regra aplicada",
    "Arquivo do razão",
    "Linha do razão",
    "Arquivo do extrato",
    "Linha ou ID do extrato",
    "Página do PDF",
    "Método de extração",
    "Parser PDF",
    "Confiança da extração",
    "Trecho do PDF",
]

DIVERGENCE_HEADERS = [
    "ID da divergência",
    "Tipo de divergência",
    "Empresa",
    "Filial",
    "Conta contábil",
    "Banco",
    "Agência",
    "Conta bancária",
    "Data do razão",
    "Documento do razão",
    "Histórico do razão",
    "Valor do razão",
    "Arquivo do razão",
    "Linha do razão",
    "Data do extrato",
    "Documento do extrato",
    "Histórico do extrato",
    "Valor do extrato",
    "Arquivo do extrato",
    "Linha ou identificador do extrato",
    "Página do PDF",
    "Método de extração",
    "Parser PDF",
    "Confiança da extração",
    "Trecho do PDF",
    "Diferença de valor",
    "Diferença em dias",
    "Tolerância permitida",
    "Regra avaliada",
    "Motivo da divergência",
    "Status da análise",
    "Justificativa do analista",
    "Ação necessária",
    "Responsável",
    "Prazo",
    "Data de regularização",
    "Conclusão",
]

LEDGER_HEADERS = [
    "ID",
    "Status",
    "Empresa",
    "Filial",
    "Conta contábil",
    "Data",
    "Documento",
    "Histórico",
    "Valor",
    "Natureza",
    "Identificador de origem",
    "Arquivo de origem",
    "Linha de origem",
    "Regra",
    "Motivo",
]

BANK_HEADERS = [
    "ID",
    "Status",
    "Banco",
    "Agência",
    "Conta bancária",
    "Data",
    "Documento",
    "Histórico",
    "Valor",
    "Natureza",
    "Identificador de origem",
    "Arquivo de origem",
    "Linha ou ID de origem",
    "Página do PDF",
    "Método de extração",
    "Parser PDF",
    "Confiança da extração",
    "Trecho do PDF",
    "Regra",
    "Motivo",
]

AMBIGUITY_HEADERS = [
    "ID",
    "Tipo",
    "ID principal",
    "Data principal",
    "Documento principal",
    "Valor principal",
    "Arquivo principal",
    "Linha principal",
    "ID candidato",
    "Data candidato",
    "Documento candidato",
    "Valor candidato",
    "Arquivo candidato",
    "Linha candidato",
    "Página do PDF candidato",
    "Trecho do PDF candidato",
    "Critério",
    "Motivo da não decisão",
]

FILE_HEADERS = [
    "Nome do arquivo",
    "Tipo",
    "Empresa",
    "Banco",
    "Agência",
    "Conta",
    "Período",
    "Quantidade de registros",
    "Total de débito",
    "Total de crédito",
    "Quantidade de páginas",
    "Método de extração",
    "Movimentações válidas",
    "Movimentações rejeitadas",
    "Quantidade de alertas",
    "Duração do processamento (s)",
    "Hash do arquivo",
    "Data e hora",
    "Situação",
]

ERROR_HEADERS = [
    "Arquivo",
    "Etapa",
    "Código do erro",
    "Descrição",
    "Registro ou linha",
    "Orientação ao usuário",
    "Data e hora",
]
