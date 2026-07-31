import json
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conciliacao.config import create_configuration_template
from conciliacao.errors import DuplicateFileError
from conciliacao.processing import ReconciliationProcessor, initialize_workspace


def configured_workspace(tmp_path: Path) -> Path:
    initialize_workspace(tmp_path)
    config_path = tmp_path / "05_Configuracao" / "CONFIGURACAO_CONTAS.xlsx"
    workbook = load_workbook(config_path)
    workbook["Contas"].append(
        [
            "SPREAD",
            "MATRIZ",
            "1.1.1.01.001",
            "Banco Itaú",
            "ITAU",
            "1234",
            "567890",
            "CSV,XLSX,OFX",
            2,
            "0,01",
            "DOC_VALOR_DATA",
            "SIM",
            "Contabilidade",
        ]
    )
    workbook.save(config_path)
    return tmp_path


def add_input_files(base: Path) -> tuple[Path, Path]:
    ledger = base / "01_Entrada" / "Razao" / "RAZAO_SPREAD_MATRIZ_202607.csv"
    statement = base / "01_Entrada" / "Extratos" / "EXTRATO_ITAU_1234_567890_202607.csv"
    ledger.write_text(
        "Empresa;Filial;Conta contábil;Data;Documento;Histórico;Débito;Crédito;Identificador\n"
        "SPREAD;MATRIZ;1.1.1.01.001;28/07/2026;001;Pagamento;100;0;L1\n"
        "SPREAD;MATRIZ;1.1.1.01.001;29/07/2026;002;Recebimento;0;250;L2\n",
        encoding="utf-8",
    )
    statement.write_text(
        "Data;Documento;Histórico;Valor;Identificador bancário\n"
        "28/07/2026;001;Pagamento;-100;B1\n"
        "29/07/2026;002;Recebimento;250;B2\n"
        "30/07/2026;003;Tarifa;-40;B3\n",
        encoding="utf-8",
    )
    return ledger, statement


def test_initialize_workspace_is_idempotent_and_does_not_overwrite_configuration(tmp_path):
    initialize_workspace(tmp_path)
    config = tmp_path / "05_Configuracao" / "CONFIGURACAO_CONTAS.xlsx"
    original = config.read_bytes()

    initialize_workspace(tmp_path)

    assert config.read_bytes() == original
    assert all(
        (tmp_path / path).is_dir()
        for path in [
            "01_Entrada/Razao",
            "01_Entrada/Extratos",
            "02_Processados",
            "03_Pendencias",
            "04_Resultados",
            "05_Configuracao",
        ]
    )


def test_complete_monthly_processing_generates_report_log_manifest_and_archives(tmp_path):
    base = configured_workspace(tmp_path)
    ledger, statement = add_input_files(base)
    original_hashes = {ledger.name: ledger.read_bytes(), statement.name: statement.read_bytes()}

    outcome = ReconciliationProcessor(base).run("202607")

    assert outcome.report_path.is_file()
    assert outcome.log_path.is_file()
    assert not ledger.exists()
    assert not statement.exists()
    archive = base / "02_Processados" / "2026-07"
    archived_ledger = archive / "Razao" / ledger.name
    archived_statement = archive / "Extratos" / statement.name
    assert archived_ledger.read_bytes() == original_hashes[ledger.name]
    assert archived_statement.read_bytes() == original_hashes[statement.name]
    manifest = json.loads((base / "02_Processados" / "processed_hashes.json").read_text())
    assert len(manifest["files"]) == 2
    log_content = outcome.log_path.read_text(encoding="utf-8")
    assert "Processamento concluído" in log_content
    assert "Duração:" in log_content
    assert "Registros válidos importados: 5" in log_content
    workbook = load_workbook(outcome.report_path, data_only=True)
    summary_headers = [cell.value for cell in workbook["Resumo"][1]]
    summary = dict(
        zip(summary_headers, [cell.value for cell in workbook["Resumo"][2]], strict=True)
    )
    assert summary["Quantidade conciliada"] == 2
    assert summary["Quantidade sem razão"] == 1


def test_same_hash_requires_explicit_reprocessing_and_never_overwrites_report(tmp_path):
    base = configured_workspace(tmp_path)
    ledger, statement = add_input_files(base)
    first = ReconciliationProcessor(base).run("202607")
    archive = base / "02_Processados" / "2026-07"
    shutil.copy2(archive / "Razao" / ledger.name, ledger)
    shutil.copy2(archive / "Extratos" / statement.name, statement)

    with pytest.raises(DuplicateFileError, match="já foi processado"):
        ReconciliationProcessor(base).run("202607")

    second = ReconciliationProcessor(base).run("202607", allow_reprocessing=True)
    assert second.report_path != first.report_path
    assert first.report_path.exists()
    assert second.report_path.exists()


def test_invalid_statement_is_sent_to_pending_while_valid_ledger_is_reported(tmp_path):
    base = configured_workspace(tmp_path)
    ledger = base / "01_Entrada" / "Razao" / "RAZAO_SPREAD_MATRIZ_202607.csv"
    statement = base / "01_Entrada" / "Extratos" / "EXTRATO_ITAU_1234_567890_202607.csv"
    ledger.write_text(
        "Empresa;Filial;Conta contábil;Data;Valor\nSPREAD;MATRIZ;1.1.1.01.001;28/07/2026;-100\n",
        encoding="utf-8",
    )
    statement.write_text("Data;Documento\n28/07/2026;001\n", encoding="utf-8")

    outcome = ReconciliationProcessor(base).run("202607")

    assert any(error.code == "COLUNAS_AUSENTES" for error in outcome.errors)
    assert (base / "03_Pendencias" / "2026-07" / statement.name).exists()
    assert (base / "02_Processados" / "2026-07" / "Razao" / ledger.name).exists()
    assert any(result.status == "RAZAO_SEM_EXTRATO" for result in outcome.results)


def test_template_helper_used_by_init_creates_a_valid_xlsx(tmp_path):
    path = create_configuration_template(tmp_path / "config.xlsx")
    assert load_workbook(path).sheetnames == ["Contas", "Layouts", "Parametros", "Instrucoes"]


def test_invalid_filename_is_archived_as_pending_and_reported(tmp_path):
    base = configured_workspace(tmp_path)
    invalid = base / "01_Entrada" / "Extratos" / "extrato com espacos.csv"
    invalid.write_text("conteúdo", encoding="utf-8")

    outcome = ReconciliationProcessor(base).run("202607")

    assert any(error.file == invalid.name for error in outcome.errors)
    assert (base / "03_Pendencias" / "2026-07" / invalid.name).exists()
