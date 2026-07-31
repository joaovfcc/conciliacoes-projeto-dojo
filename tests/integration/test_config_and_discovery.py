from decimal import Decimal

import pytest
from openpyxl import load_workbook

from conciliacao.config import create_configuration_template, load_configuration
from conciliacao.discovery import discover_input_files, file_sha256, parse_input_filename
from conciliacao.errors import ConfigurationError, FileValidationError


def test_configuration_template_round_trip(tmp_path):
    config_path = tmp_path / "05_Configuracao" / "CONFIGURACAO_CONTAS.xlsx"
    create_configuration_template(config_path)
    workbook = load_workbook(config_path)
    accounts = workbook["Contas"]
    accounts.append(
        [
            "SPREAD",
            "MATRIZ",
            "1.1.1.01.001",
            "Banco Itaú",
            "ITAU",
            "1234",
            "567890",
            "OFX,CSV,XLSX",
            2,
            "0,01",
            "DOC_VALOR_DATA",
            "SIM",
            "Contabilidade",
        ]
    )
    workbook.save(config_path)

    config = load_configuration(config_path)

    assert len(config.accounts) == 1
    assert config.accounts[0].bank_key == ("ITAU", "1234", "567890")
    assert config.accounts[0].value_tolerance == Decimal("0.01")
    assert config.parameters["PERIODO_PADRAO"] == ""
    assert len(config.sha256) == 64


def test_filename_parser_accepts_contract_and_rejects_unsafe_name(tmp_path):
    ledger = parse_input_filename(tmp_path / "RAZAO_SPREAD_MATRIZ_202607.xlsx")
    bank = parse_input_filename(tmp_path / "EXTRATO_ITAU_1234_567890_202607.ofx")

    assert (ledger.file_type, ledger.company, ledger.period) == ("RAZAO", "SPREAD", "202607")
    assert (bank.bank, bank.agency, bank.account) == ("ITAU", "1234", "567890")

    with pytest.raises(FileValidationError, match="nomenclatura"):
        parse_input_filename(tmp_path / "extrato itau julho.csv")


def test_discovery_ignores_temporary_files_and_detects_content_hash(tmp_path):
    ledger_dir = tmp_path / "01_Entrada" / "Razao"
    bank_dir = tmp_path / "01_Entrada" / "Extratos"
    ledger_dir.mkdir(parents=True)
    bank_dir.mkdir(parents=True)
    ledger_file = ledger_dir / "RAZAO_SPREAD_MATRIZ_202607.csv"
    bank_file = bank_dir / "EXTRATO_ITAU_1234_567890_202607.csv"
    temporary = bank_dir / "~$EXTRATO_ITAU_1234_567890_202607.xlsx"
    ledger_file.write_text("razao", encoding="utf-8")
    bank_file.write_text("extrato", encoding="utf-8")
    temporary.write_text("lock", encoding="utf-8")

    discovered, errors = discover_input_files(tmp_path, "202607")

    assert not errors
    assert {item.path for item in discovered} == {ledger_file, bank_file}
    assert all(len(item.sha256) == 64 for item in discovered)
    assert file_sha256(ledger_file) == next(
        item.sha256 for item in discovered if item.path == ledger_file
    )


def test_discovery_reports_wrong_period_without_silently_processing(tmp_path):
    folder = tmp_path / "01_Entrada" / "Razao"
    folder.mkdir(parents=True)
    path = folder / "RAZAO_SPREAD_MATRIZ_202606.csv"
    path.write_text("content", encoding="utf-8")

    discovered, errors = discover_input_files(tmp_path, "202607")

    assert discovered == []
    assert errors[0].code == "PERIODO_DIVERGENTE"
    assert errors[0].file == path.name


def test_configuration_rejects_duplicate_account_relationships(tmp_path):
    config_path = tmp_path / "CONFIGURACAO_CONTAS.xlsx"
    create_configuration_template(config_path)
    workbook = load_workbook(config_path)
    row = [
        "SPREAD",
        "MATRIZ",
        "1.1.1.01.001",
        "Banco Itaú",
        "ITAU",
        "1234",
        "567890",
        "OFX",
        2,
        "0,01",
        "DOC_VALOR_DATA",
        "SIM",
        "",
    ]
    workbook["Contas"].append(row)
    workbook["Contas"].append(row)
    workbook.save(config_path)

    with pytest.raises(ConfigurationError, match="duplicado"):
        load_configuration(config_path)


def test_discovery_rejects_duplicate_content_in_the_same_execution(tmp_path):
    folder = tmp_path / "01_Entrada" / "Extratos"
    folder.mkdir(parents=True)
    first = folder / "EXTRATO_ITAU_1234_567890_202607.csv"
    second = folder / "EXTRATO_ITAU_1234_999999_202607.csv"
    first.write_text("mesmo conteúdo", encoding="utf-8")
    second.write_text("mesmo conteúdo", encoding="utf-8")

    discovered, errors = discover_input_files(tmp_path, "202607")

    assert len(discovered) == 1
    assert errors[0].code == "ARQUIVO_DUPLICADO"
    assert discovered[0].path == first
